"""Step 12: Build the client-facing Excel workbook.

Produces a 7-tab workbook that reads like a local managed-IT opportunity map:
Executive Summary, Final Sales Ready Accounts, Account Briefings, Contact
Details, Verification Sources, Backup Accounts, and QA Report. Only client-safe
language is written; no credentials, scores, or methodology terms appear.
"""
import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as c

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
TIER_FILLS = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="FFEB9C"),
    "C": PatternFill("solid", fgColor="FCE4D6"),
    "D": PatternFill("solid", fgColor="F2F2F2"),
}


def _week_label(week):
    n = "".join(ch for ch in week if ch.isdigit())
    return f"Week{n}" if n else week.title().replace("_", "")


def _style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _exec_summary(wb, final, backup, week, qa):
    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = f"CMIT Solutions of Glendale - Sales Ready Opportunity Map ({_week_label(week)})"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    ws["A3"] = "A curated, local managed IT opportunity map for LA County SMBs."
    ws["A3"].font = Font(italic=True, size=11)

    tiers = {}
    verts = {}
    for a in final:
        tiers[a.get("priority_tier_label", "")] = tiers.get(a.get("priority_tier_label", ""), 0) + 1
        verts[a.get("vertical", "")] = verts.get(a.get("vertical", ""), 0) + 1

    rows = [
        ("", ""),
        ("Total sales-ready accounts", len(final)),
        ("Verified decision makers", sum(1 for a in final if a.get("verified_email"))),
        ("Backup accounts available", len(backup)),
        ("QA status", qa.get("status", "not run").upper() if qa else "NOT RUN"),
        ("", ""),
        ("Accounts by priority tier", ""),
    ]
    for label, n in sorted(tiers.items()):
        rows.append((f"   {label}", n))
    rows.append(("", ""))
    rows.append(("Accounts by vertical", ""))
    for label, n in sorted(verts.items(), key=lambda x: -x[1]):
        rows.append((f"   {label}", n))

    r = 5
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=bool(label and not label.startswith(" ")))
        ws.cell(row=r, column=2, value=val)
        r += 1
    _autosize(ws, [40, 20])


def _final_tab(wb, final):
    ws = wb.create_sheet("Final Sales Ready Accounts")
    headers = [
        "Rank", "Sales Priority Tier", "Company", "Vertical", "Sub Vertical",
        "Website", "Domain", "Address", "City", "Estimated Employees",
        "Employee Confidence", "Decision Maker", "Title", "Verified Email",
        "Direct Phone", "Company Phone", "LinkedIn", "Managed IT Trigger",
        "Specific Visible Risk", "Compliance Angle", "Why CMIT Should Call",
        "Recommended Opening Line", "Suggested Discovery Question", "Source Evidence",
    ]
    ws.append(headers)
    for rank, a in enumerate(final, 1):
        ws.append([
            rank, a.get("priority_tier_label", ""), a.get("company_name", ""),
            a.get("vertical", ""), a.get("sub_vertical", ""), a.get("website", ""),
            a.get("domain", ""), a.get("address", ""), a.get("city", ""),
            a.get("estimated_employees_display", ""), a.get("employee_confidence", ""),
            a.get("decision_maker", ""), a.get("decision_maker_title", ""),
            a.get("verified_email", ""), a.get("direct_phone", ""), a.get("main_phone", ""),
            a.get("contact_linkedin", "") or a.get("linkedin_company_url", ""),
            a.get("managed_it_trigger", ""), a.get("specific_visible_risk", ""),
            a.get("compliance_angle", ""), a.get("why_cmit_should_call", ""),
            a.get("recommended_opening_line", ""), a.get("suggested_discovery_question", ""),
            a.get("source_evidence", ""),
        ])
        tier = a.get("priority_tier", "D")
        ws.cell(row=rank + 1, column=2).fill = TIER_FILLS.get(tier, TIER_FILLS["D"])
    _style_header(ws)
    widths = [6, 20, 26, 16, 18, 28, 22, 30, 14, 14, 12, 20, 22, 28, 16, 16, 28,
              45, 32, 32, 55, 55, 55, 60]
    _autosize(ws, widths)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _briefings_tab(wb, final):
    ws = wb.create_sheet("Account Briefings")
    for rank, a in enumerate(final, 1):
        base = ws.max_row + (2 if rank > 1 else 0)
        ws.cell(row=base + 1, column=1, value=f"{rank}. {a.get('company_name','')} - {a.get('sub_vertical','')}, {a.get('city','')}").font = Font(bold=True, size=13, color="1F4E78")
        lines = [
            ("Priority", a.get("priority_tier_label", "")),
            ("Estimated size", f"{a.get('estimated_employees_display','')} employees ({a.get('employee_confidence','')} confidence)"),
            ("Decision maker", f"{a.get('decision_maker','')} - {a.get('decision_maker_title','')}"),
            ("Managed IT trigger", a.get("managed_it_trigger", "")),
            ("Specific visible risk", a.get("specific_visible_risk", "")),
            ("Compliance angle", a.get("compliance_angle", "")),
            ("Security posture", a.get("security_posture_summary", "")),
            ("Email security", a.get("email_security_summary", "")),
            ("Website security", a.get("website_security_summary", "")),
            ("Why CMIT should call", a.get("why_cmit_should_call", "")),
            ("Recommended opening line", a.get("recommended_opening_line", "")),
            ("Suggested discovery question", a.get("suggested_discovery_question", "")),
        ]
        r = base + 2
        for label, val in lines:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            cell = ws.cell(row=r, column=2, value=val)
            cell.alignment = WRAP
            r += 1
    _autosize(ws, [28, 110])


def _contacts_tab(wb, final):
    ws = wb.create_sheet("Contact Details")
    headers = ["Rank", "Company", "Decision Maker", "Title", "Verified Email",
               "Email Status", "Direct Phone", "Company Phone", "LinkedIn", "City", "Domain"]
    ws.append(headers)
    for rank, a in enumerate(final, 1):
        ws.append([
            rank, a.get("company_name", ""), a.get("decision_maker", ""),
            a.get("decision_maker_title", ""), a.get("verified_email", ""),
            a.get("email_status", ""), a.get("direct_phone", ""), a.get("main_phone", ""),
            a.get("contact_linkedin", ""), a.get("city", ""), a.get("domain", ""),
        ])
    _style_header(ws)
    _autosize(ws, [6, 26, 22, 22, 30, 14, 16, 16, 32, 14, 22])


def _sources_tab(wb, final):
    ws = wb.create_sheet("Verification Sources")
    headers = ["Rank", "Company", "Website", "Business Listing", "Company LinkedIn",
               "Estimated Size Evidence", "Source Evidence"]
    ws.append(headers)
    for rank, a in enumerate(final, 1):
        ws.append([
            rank, a.get("company_name", ""), a.get("website", ""),
            a.get("google_maps_url", ""), a.get("linkedin_company_url", ""),
            a.get("employee_evidence", ""), a.get("source_evidence", ""),
        ])
    _style_header(ws)
    _autosize(ws, [6, 26, 30, 34, 32, 50, 60])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _backup_tab(wb, backup):
    ws = wb.create_sheet("Backup Accounts")
    headers = ["Company", "Vertical", "Sub Vertical", "City", "Website", "Domain",
               "Estimated Employees", "Company Phone", "Managed IT Trigger", "Status"]
    ws.append(headers)
    for a in backup[:200]:
        ws.append([
            a.get("company_name", ""), a.get("vertical", ""), a.get("sub_vertical", ""),
            a.get("city", ""), a.get("website", ""), a.get("domain", ""),
            a.get("estimated_employees_display", ""), a.get("main_phone", ""),
            a.get("managed_it_trigger", ""), a.get("_backup_reason", ""),
        ])
    _style_header(ws)
    _autosize(ws, [26, 16, 18, 14, 28, 22, 16, 16, 45, 22])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _qa_tab(wb, qa):
    ws = wb.create_sheet("QA Report")
    ws["A1"] = "QA Report"
    ws["A1"].font = TITLE_FONT
    if not qa:
        ws["A3"] = "QA has not been run yet."
        _autosize(ws, [50, 20])
        return
    ws.cell(row=3, column=1, value="Overall status").font = Font(bold=True)
    status_cell = ws.cell(row=3, column=2, value=qa.get("status", "").upper())
    status_cell.fill = TIER_FILLS["A"] if qa.get("status") == "pass" else PatternFill("solid", fgColor="FFC7CE")
    ws.cell(row=5, column=1, value="Check").font = HEADER_FONT
    ws.cell(row=5, column=2, value="Result").font = HEADER_FONT
    ws.cell(row=5, column=3, value="Detail").font = HEADER_FONT
    for cell in ws[5]:
        if cell.value:
            cell.fill = HEADER_FILL
    r = 6
    for chk in qa.get("checks", []):
        ws.cell(row=r, column=1, value=chk["name"])
        rc = ws.cell(row=r, column=2, value="PASS" if chk["passed"] else "FAIL")
        rc.fill = TIER_FILLS["A"] if chk["passed"] else PatternFill("solid", fgColor="FFC7CE")
        ws.cell(row=r, column=3, value=str(chk.get("detail", ""))).alignment = WRAP
        r += 1
    _autosize(ws, [42, 12, 60])


def build(week):
    log = c.get_logger("build_workbook", week)
    final = c.load_json(c.processed_path(week, "accounts_final.json"), []) or []
    backup = c.load_json(c.processed_path(week, "accounts_backup.json"), []) or []
    qa = c.load_json(c.processed_path(week, "qa_results.json"), None)

    wb = Workbook()
    _exec_summary(wb, final, backup, week, qa)
    _final_tab(wb, final)
    _briefings_tab(wb, final)
    _contacts_tab(wb, final)
    _sources_tab(wb, final)
    _backup_tab(wb, backup)
    _qa_tab(wb, qa)

    out = c.week_dir(week) / f"CMIT_Glendale_Sales_Ready_Database_{_week_label(week)}.xlsx"
    wb.save(out)
    log.info("workbook written -> %s (%d final, %d backup)", out, len(final), len(backup))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--week", required=True)
    args = ap.parse_args()
    c.load_env()
    build(args.week)
